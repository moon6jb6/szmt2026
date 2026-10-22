from pptx import Presentation
from pptx.util import Cm
p=r'D:\提交\作品汇报PPT\汇报PPT_把名字留下来_460737.pptx'
prs=Presentation(p)
def all_shapes(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type==6:
            yield from all_shapes(sh.shapes)
# 16 页结语上移，避开册子图
slide=prs.slides[15]
for sh in all_shapes(slide.shapes):
    if sh.name=='L2-结语' and sh.has_text_frame:
        sh.top=Cm(7.0)
        print('moved slide16 结语 to 7.0cm')
# 13 页构建体积改当前值
pairs=[('298.5 kB','316.6 kB'),('23.47 MB','24.58 MB')]
count=0
for sh in all_shapes(prs.slides[12].shapes):
    if not getattr(sh,'has_text_frame',False): continue
    for para in sh.text_frame.paragraphs:
        full=''.join(r.text for r in para.runs)
        new=full
        for a,b in pairs: new=new.replace(a,b)
        if new!=full and para.runs:
            para.runs[0].text=new
            for r in para.runs[1:]: r.text=''
            count+=1
print('number fixes',count)
prs.save(p)
print('saved')
