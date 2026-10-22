# -*- coding: utf-8 -*-
"""检查汇报 PPT 的结构：每页文本、图片、空框（候选截图位）。"""
from pptx import Presentation
from pptx.util import Emu

PPT = r'D:\提交\作品汇报PPT\汇报PPT_把名字留下来_460737.pptx'
OUT = r'D:\backups\ppt-inspect.txt'

prs = Presentation(PPT)
lines = []
lines.append(f'slides={len(prs.slides)} size={Emu(prs.slide_width).cm:.1f}x{Emu(prs.slide_height).cm:.1f}cm')


def cm(v):
    return Emu(v).cm


def walk(shapes, depth, acc, empties, pictures, texts):
    for sh in shapes:
        pad = '  ' * depth
        st = sh.shape_type
        try:
            left, top, w, h = cm(sh.left), cm(sh.top), cm(sh.width), cm(sh.height)
        except Exception:
            left = top = w = h = -1
        if st == 6:  # group
            kids = list(sh.shapes)
            acc.append(f'{pad}GROUP {sh.name!r} ({left:.1f},{top:.1f}) {w:.1f}x{h:.1f} kids={len(kids)}')
            walk(kids, depth + 1, acc, empties, pictures, texts)
            continue
        txt = ''
        try:
            if sh.has_text_frame:
                txt = sh.text_frame.text.replace('\n', ' / ')[:60]
        except Exception:
            pass
        is_pic = False
        try:
            is_pic = sh.shape_type == 13
        except Exception:
            pass
        if is_pic:
            pictures.append(f'{pad}PIC {sh.name!r} ({left:.1f},{top:.1f}) {w:.1f}x{h:.1f}')
        elif txt:
            texts.append(txt)
        elif w > 4 and h > 2.5:
            empties.append(f'{pad}EMPTY {sh.name!r} ({left:.1f},{top:.1f}) {w:.1f}x{h:.1f}')
        acc.append(f'{pad}{st} {sh.name!r} ({left:.1f},{top:.1f}) {w:.1f}x{h:.1f} txt={txt!r}')


for idx, slide in enumerate(prs.slides, 1):
    acc, empties, pictures, texts = [], [], [], []
    walk(slide.shapes, 0, acc, empties, pictures, texts)
    title = texts[0] if texts else ''
    lines.append(f'\n=== slide {idx}  {title}')
    for t in texts:
        lines.append(f'  TXT: {t}')
    for p in pictures:
        lines.append(f'  {p}')
    for e in empties:
        lines.append(f'  {e}')
    # 完整树收进 details
    lines.append('  --- tree ---')
    lines.extend(acc)

open(OUT, 'w', encoding='utf-8').write('\n'.join(lines))
print('written', OUT, 'lines', len(lines))
