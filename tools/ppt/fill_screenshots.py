# -*- coding: utf-8 -*-
"""把运行截图填进汇报 PPT 的 7 个占位框，另存为 *_含截图.pptx。"""
from pathlib import Path

from docx.shared import Cm
from pptx import Presentation
from PIL import Image

SRC = Path(r'D:\提交\作品汇报PPT\汇报PPT_把名字留下来_460737.pptx')
DST = Path(r'D:\提交\作品汇报PPT\汇报PPT_把名字留下来_460737_含截图.pptx')
IMG = Path(r'D:\提交\截图素材')

# 每个占位：删除这些名字的形状，然后在指定位置贴图（单位 cm）
SLOTS = [
    {
        'slide': 7,
        'remove': ['L3-截图占位框', 'L3-占位提示', 'L3-占位文件名', 'L3-占位尺寸'],
        'boxes': [(2.2, 6.2, 24.9, 14.0)],
        'images': ['slide7-2x-v0-03.png'],
    },
    {
        'slide': 8,
        'remove': ['L3-截图占位框', 'L3-占位提示', 'L3-占位文件名', 'L3-占位尺寸'],
        'boxes': [(17.9, 6.0, 6.7, 3.8), (25.4, 6.0, 6.7, 3.8), (17.9, 11.4, 6.7, 3.8), (25.4, 11.4, 6.7, 3.8)],
        'images': ['ink-cover.png', 'ink-hold.png', 'ink-reveal.png', 'finale-帧4-5.6s.png'],
    },
    {
        'slide': 9,
        'remove': ['L3-截图占位框', 'L3-占位提示', 'L3-占位文件名', 'L3-占位尺寸'],
        'boxes': [(18.4, 7.4, 8.9, 5.0)],
        'images': ['3d-v2-44-桥.png'],
    },
    {
        'slide': 15,
        'remove': ['L3-截图占位框', 'L3-占位提示', 'L3-占位文件名', 'L3-占位尺寸'],
        'boxes': [(24.3, 3.9, 6.0, 13.1)],
        'images': ['mobile-桥头.png'],
    },
]


def all_shapes(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:  # group
            yield from all_shapes(sh.shapes)


def remove_named(slide, names):
    targets = []
    for sh in all_shapes(slide.shapes):
        if sh.name in names:
            targets.append(sh)
    for sh in targets:
        el = sh._element
        el.getparent().remove(el)
    return len(targets)


def add_cropped(slide, image_path, left, top, width, height):
    pic = slide.shapes.add_picture(str(image_path), Cm(left), Cm(top), Cm(width), Cm(height))
    im = Image.open(image_path)
    img_ar = im.width / im.height
    box_ar = width / height
    if img_ar > box_ar:
        keep = box_ar / img_ar
        crop = (1 - keep) / 2
        pic.crop_left = crop
        pic.crop_right = crop
    elif img_ar < box_ar:
        keep = img_ar / box_ar
        crop = (1 - keep) / 2
        pic.crop_top = crop
        pic.crop_bottom = crop
    return pic


def main():
    prs = Presentation(str(SRC))
    for slot in SLOTS:
        slide = prs.slides[slot['slide'] - 1]
        removed = remove_named(slide, slot['remove'])
        for (left, top, width, height), name in zip(slot['boxes'], slot['images']):
            path = IMG / name
            if not path.exists():
                raise SystemExit(f'missing image: {path}')
            add_cropped(slide, path, left, top, width, height)
        print(f'slide {slot["slide"]}: removed {removed} placeholder shapes, added {len(slot["images"])} images')
    prs.save(str(DST))
    print('saved', DST, DST.stat().st_size, 'bytes')


if __name__ == '__main__':
    main()
