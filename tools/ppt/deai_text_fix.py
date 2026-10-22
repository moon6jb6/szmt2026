# -*- coding: utf-8 -*-
"""补齐 deai_text.py 漏掉的段落：同一段落里含多个句子、或跨 run 的情况。"""
from pathlib import Path

from pptx import Presentation

PPT = Path(r'D:\提交\作品汇报PPT\汇报PPT_把名字留下来_460737.pptx')

PAIRS = [
    ('写得多不难。难的是，每一句都敢让人查。', '写够字数不难。难的是每一处史实都能查到出处。'),
    # 兜底：如果上面这条形状里是单段，直接整段换
    ('写得多不难。', '写够字数不难。'),
    ('难的是，每一句都敢让人查。', '难的是每一处史实都能查到出处。'),
]


def all_shapes(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from all_shapes(sh.shapes)


def replace_paragraph(para, old, new):
    full = ''.join(run.text for run in para.runs)
    if old not in full:
        return False
    if full == old:
        runs = para.runs
        if runs:
            runs[0].text = new
            for run in runs[1:]:
                run.text = ''
        return True
    # 单 run 内替换
    for run in para.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            return True
    # 跨 run：压缩到第一个 run
    runs = para.runs
    if runs:
        runs[0].text = full.replace(old, new)
        for run in runs[1:]:
            run.text = ''
        return True
    return False


prs = Presentation(str(PPT))
count = 0
for slide in prs.slides:
    for shape in all_shapes(slide.shapes):
        if not getattr(shape, 'has_text_frame', False):
            continue
        for para in shape.text_frame.paragraphs:
            for old, new in PAIRS:
                if replace_paragraph(para, old, new):
                    count += 1
                    break
prs.save(str(PPT))
print('extra replaced', count)
