#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Markdown → DOCX（把交付文档 md 生成排版干净的 Word；用于 AI 工具使用情况说明等）。

用法：
    python tools/docgen/md-to-docx.py docs/deliverables/人工智能工具使用情况说明.md
    python tools/docgen/md-to-docx.py <in.md> <out.docx>
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]


def inline_runs(p, text):
    """处理 **加粗**、`code`、[文字](链接)，其余按普通文字。"""
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            p.add_run(text[pos:m.start()])
        token = m.group(0)
        if token.startswith('**'):
            r = p.add_run(token[2:-2])
            r.bold = True
        elif token.startswith('`'):
            r = p.add_run(token[1:-1])
            r.font.name = 'Consolas'
        else:
            label = token[1:token.index(']')]
            r = p.add_run(label)
            r.font.color.rgb = RGBColor(0x7A, 0x3A, 0x2A)
        pos = m.end()
    if pos < len(text):
        p.add_run(text[pos:])


def build(md_path, out_path):
    text = Path(md_path).read_text(encoding='utf-8-sig').replace('\r\n', '\n')
    lines = text.split('\n')
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'SimSun'
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(
        '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia', 'SimSun'
    )
    ink = RGBColor(0x1A, 0x1A, 0x1A)
    i = 0

    def flush_para(buf):
        if not buf:
            return
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.4
        p.paragraph_format.space_after = Pt(4)
        inline_runs(p, ''.join(buf))
        for run in p.runs:
            run.font.name = 'SimSun'
            run.font.size = Pt(10.5)
            run.font.color.rgb = ink

    while i < len(lines):
        line = lines[i]
        if line.startswith('```'):
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                code.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            r = p.add_run('\n'.join(code))
            r.font.name = 'Consolas'
            r.font.size = Pt(9)
            continue
        m = re.match(r'^(#{1,4})\s+(.*)$', line)
        if m:
            level = len(m.group(1))
            p = doc.add_paragraph()
            r = p.add_run(m.group(2).strip())
            r.bold = True
            r.font.name = 'SimSun'
            r.font.color.rgb = ink
            r.font.size = Pt(18 if level == 1 else 14 if level == 2 else 12)
            if level == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
            p.paragraph_format.space_after = Pt(6)
            i += 1
            continue
        if re.match(r'^\s*\|', line):
            rows = []
            while i < len(lines) and re.match(r'^\s*\|', lines[i]):
                rows.append(lines[i].strip())
                i += 1
            cells = [r.strip().strip('|').split('|') for r in rows]
            cells = [[c.strip() for c in row] for row in cells]
            body = [row for row in cells if not all(re.match(r'^:?-{2,}:?$', c) for c in row)]
            if not body:
                continue
            t = doc.add_table(rows=1, cols=len(body[0]))
            t.style = 'Table Grid'
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, cell in enumerate(body[0]):
                par = t.rows[0].cells[j].paragraphs[0]
                r = par.add_run(re.sub(r'\*\*([^*]+)\*\*', r'\1', cell))
                r.bold = True
                r.font.name = 'SimSun'
                r.font.size = Pt(10)
            for row in body[1:]:
                cells_row = t.add_row().cells
                for j, cell in enumerate(row[:len(cells_row)]):
                    par = cells_row[j].paragraphs[0]
                    inline_runs(par, re.sub(r'\*\*([^*]+)\*\*', r'\1', cell))
                    for run in par.runs:
                        run.font.name = 'SimSun'
                        run.font.size = Pt(9.5)
            doc.add_paragraph()
            continue
        if re.match(r'^\s*[-*]\s+', line):
            while i < len(lines) and re.match(r'^\s*[-*]\s+', lines[i]):
                p = doc.add_paragraph(style='List Bullet')
                inline_runs(p, re.sub(r'^\s*[-*]\s+', '', lines[i]))
                for run in p.runs:
                    run.font.name = 'SimSun'
                    run.font.size = Pt(10.5)
                i += 1
            continue
        if re.match(r'^\s*\d+[.)]\s+', line):
            while i < len(lines) and re.match(r'^\s*\d+[.)]\s+', lines[i]):
                p = doc.add_paragraph(style='List Number')
                inline_runs(p, re.sub(r'^\s*\d+[.)]\s+', '', lines[i]))
                for run in p.runs:
                    run.font.name = 'SimSun'
                    run.font.size = Pt(10.5)
                i += 1
            continue
        if re.match(r'^\s*>\s?', line):
            buf = []
            while i < len(lines) and re.match(r'^\s*>\s?', lines[i]):
                buf.append(re.sub(r'^\s*>\s?', '', lines[i]))
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            inline_runs(p, '\n'.join(buf))
            for run in p.runs:
                run.font.name = 'SimSun'
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(0x3B, 0x33, 0x2A)
            continue
        if re.match(r'^(-{3,}|\*{3,})\s*$', line):
            doc.add_paragraph('─' * 40)
            i += 1
            continue
        if line.strip() == '':
            i += 1
            continue
        buf = [line]
        i += 1
        while i < len(lines) and lines[i].strip() != '' and not re.match(
            r'^(#{1,4}\s|\s*\||\s*[-*]\s+|\s*\d+[.)]\s+|>\s?|```)', lines[i]
        ):
            buf.append(lines[i])
            i += 1
        flush_para(buf)

    doc.save(out_path)
    print(f'saved: {out_path} ({Path(out_path).stat().st_size} bytes)')


if __name__ == '__main__':
    src = sys.argv[1] if len(sys.argv) > 1 else 'docs/deliverables/人工智能工具使用情况说明.md'
    dst = sys.argv[2] if len(sys.argv) > 2 else str(Path(src).with_name(
        Path(src).stem + '_460737.docx'
    ))
    build(Path(src), dst)
